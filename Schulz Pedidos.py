import tkinter as tk
from tkinter import simpledialog, messagebox, filedialog
from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from selenium.webdriver.common.keys import Keys
from selenium.webdriver.support.ui import Select 
from selenium.common.exceptions import TimeoutException, StaleElementReferenceException, ElementClickInterceptedException
import time
import os 
import pandas as pd
from dotenv import load_dotenv 
import re
import json
from openpyxl.utils import get_column_letter, column_index_from_string

# --- CONFIGURAÇÃO PERSISTENTE (TABELA SCHULZ E COLUNA) ---
CONFIG_FILE = os.path.join(os.path.dirname(os.path.abspath(__file__)), "config_schulz.json")

def carregar_config():
    config_padrao = {
        "tabela_schulz_path": "",
        "coluna_preco": "I"
    }
    if os.path.exists(CONFIG_FILE):
        try:
            with open(CONFIG_FILE, "r", encoding="utf-8") as f:
                dados = json.load(f)
                config_padrao.update(dados)
        except Exception:
            pass
    # Tenta localizar automaticamente uma tabela Schulz na pasta se ainda não configurada
    if not config_padrao["tabela_schulz_path"] or not os.path.exists(config_padrao["tabela_schulz_path"]):
        dir_atual = os.path.dirname(os.path.abspath(__file__))
        for fname in os.listdir(dir_atual):
            if "Tabela Schulz" in fname and fname.endswith((".xlsx", ".xls")):
                config_padrao["tabela_schulz_path"] = os.path.join(dir_atual, fname)
                salvar_config(config_padrao)
                break
    return config_padrao

def salvar_config(config_dict):
    try:
        with open(CONFIG_FILE, "w", encoding="utf-8") as f:
            json.dump(config_dict, f, indent=4, ensure_ascii=False)
    except Exception:
        pass

# --- 1. CONFIGURAÇÕES ---

load_dotenv()

URL_PORTAL_LF = "https://erp.bem.inf.br/ords/f?p=8000000:9999"

# Dados Fixos Schulz Unificado
CNPJS_SCHULZ = {
    "Automotiva": "84693183000168",
    "Compressores": "23635798000143"
}
CONDICAO_PAGTO_TEXTO = "90" 
FRETE_TIPO = "CIF"
CNPJ_TRANSPORTADORA = "53417977000175" # JEVERSON MELO

# --- TEXTOS DE OBSERVAÇÃO ---
# Texto que vai no campo "Observação NF" (P7_MSG_OBS) quando é Industrialização 
OBS_NF_INDUSTRIALIZACAO = "Diferimento de ICMS TTD 490\nTermo de concessão nº 205000004022828"

# Texto FIXO que vai no campo "Observação Pedido" (P7_OBS)
OBS_FIXA_PEDIDO = "Preço autorizado Rafael Schulze"

CREDENCIAIS = {
    "LF": {
        "user": os.environ.get("LF_USER"), 
        "pass": os.environ.get("LF_PASS")
    }
}

if not all([CREDENCIAIS["LF"]["user"], CREDENCIAIS["LF"]["pass"]]):
    messagebox.showerror("Erro de Credenciais", "Configure as variáveis de ambiente LF_USER e LF_PASS no arquivo .env")
    exit()

# XPATHs LF 
XPATHS_LF = {
    "login_user": '//*[@id="P9999_USERNAME"]',
    "login_pass": '//*[@id="P9999_PASSWORD"]',
    "login_btn": "//button[contains(., 'Acessar')]", 
    "nav_menu_btn": "/html/body/form/header/div[1]/div[1]/button",
    "nav_pedido_link": "/html/body/form/div[1]/div[1]/div/ul/li[9]/div[2]/a",
    "novo_pedido_btn_opcoes": [
        "//button[span[contains(text(), 'Novo Pedido')]]",
        "//button[contains(., 'Novo Pedido')]",
        "//button[contains(., 'Novo')]",
        "/html/body/form/div[1]/div[2]/div[2]/main/div[2]/div/div[1]/div/div/div[1]/div[3]/button[1]"
    ],
    "novo_pedido_iframe": "//iframe[@title='Pedido']", 
    "novo_pedido_iframe_check": "P7_AD_ORG_ID", 
    "campo_filial": "P7_AD_ORG_ID",
    "campo_finalidade": "P7_BEM_VEN_FINALIDADE_ID",
    "campo_tipo_nf": "P7_BEG_TIPO_NF_ID",
    "campo_ordem_compra": "P7_NRO_ORDEM_COMPRA",
    "campo_tipo_frete": "P7_BEM_TIPO_FRETE_ID",
    
    # --- CAMPOS DE OBSERVAÇÃO ---
    "campo_obs_nf": "P7_MSG_OBS",      # Onde vai o TTD 490
    "campo_obs_pedido": "P7_OBS",      # Onde vai "Preço autorizado..."
    # ----------------------------

    "btn_salvar_cabecalho": "//button[span[normalize-space(.)='Incluir'] and not(ancestor::div[contains(@class,'ui-dialog-content')])]",
    "btn_adicionar_item": "//button[@title='Novo Item Pedido']", 
    "item_popup_iframe": "//iframe[@title='Itens do Pedido']",
    "item_popup_iframe_check": "P8_BEG_PRD_SKU_ID_lov_btn", 
    "item_popup_sku_btn": "P8_BEG_PRD_SKU_ID_lov_btn",
    "item_sku_search_result": "//div[contains(@class, 'a-PopupLOV-results')]//table/tbody/tr[1]/td[1]",
    "item_popup_ordem_compra": "P8_NRO_ORDEM_COMPRA", 
    "item_popup_item_oc": "P8_NRO_ITEM_OC", 
    "item_popup_qtd": "P8_QUANTIDADE",
    "item_popup_preco": "P8_VLR_UNIT_SUGERIDO",
    "item_popup_btn_salvar": "btn-incluir",
    
    # Radio
    "radio_ped_complementar_sim_label": "//label[@for='P7_IND_PED_COMPLEMENTAR_0']",

    # Botões do Cabeçalho para abrir Popups (Centralizados)
    "btn_cliente_lov": "//button[contains(@id,'P7_BEG_VW_CLIENTE_ID_lov_btn')]",
    "btn_forma_pgto_lov": "//button[contains(@id,'P7_BEG_VEN_FORMA_PAGAMENTO_ID_lov_btn')]",
    "btn_condicao_pgto_lov": "//button[contains(@id,'P7_BEG_CONDICAO_PAGAMENTO_ID_lov_btn')]",
    "btn_transportadora_lov": "//button[contains(@id,'P7_BEG_VW_TRANSPORTADORA_ID_lov_btn')]"
}

# --- FUNÇÕES ROBUSTAS ---

def dismiss_apex_alert_if_any(driver, log_func=None):
    """Fecha automaticamente alertas ou popups de erro de banco do ERP LF (ex: ORA-00904 BEG_VW_TRANSP_PREF_ID)."""
    if not driver:
        return False
    dismissed = False
        
    # 1. Alerta nativo do navegador (window.alert)
    try:
        alert = driver.switch_to.alert
        texto = alert.text
        alert.accept()
        msg = f"   ℹ️ [ERP] Alerta do sistema confirmado/dispensado: '{texto}'"
        if log_func: log_func(msg)
        else: print(msg)
        dismissed = True
    except Exception:
        pass

    # 2. Executa script JavaScript que varre o iframe atual e window.top (em caso de foco cruzado)
    try:
        js_script = """
        try {
            var dismissed = false;
            var docs = [document];
            try {
                if (window.top && window.top.document && window.top.document !== document) {
                    docs.push(window.top.document);
                }
            } catch(e) {}

            for (var i = 0; i < docs.length; i++) {
                var doc = docs[i];
                var dialogs = doc.querySelectorAll(".ui-dialog:not(.ui-dialog-popuplov)");
                for (var j = 0; j < dialogs.length; j++) {
                    var d = dialogs[j];
                    if (d.offsetParent !== null) {
                        var txt = (d.innerText || "").toLowerCase();
                        if (txt.indexOf("beg_vw_transp_pref_id") !== -1 || 
                            txt.indexOf("identificador") !== -1 || 
                            txt.indexOf("ora-") !== -1 || 
                            txt.indexOf("erro") !== -1 ||
                            txt.indexOf("transportadora") !== -1) {
                            var btns = d.querySelectorAll(".ui-dialog-buttonpane button, button.js-confirm, button");
                            for (var k = 0; k < btns.length; k++) {
                                btns[k].click();
                                dismissed = true;
                                break;
                            }
                        }
                    }
                }
                var alerts = doc.querySelectorAll("#apexPageMessage, .a-AlertMessage, .t-Alert--error, .t-Alert");
                for (var a = 0; a < alerts.length; a++) {
                    var el = alerts[a];
                    if (el.offsetParent !== null) {
                        var btns = el.querySelectorAll("button");
                        for (var b = 0; b < btns.length; b++) {
                            btns[b].click();
                            dismissed = true;
                        }
                    }
                }
            }
            return dismissed;
        } catch(e) {
            return false;
        }
        """
        js_res = driver.execute_script(js_script)
        if js_res:
            msg = "   ℹ️ [ERP] Popup de aviso do ERP fechado automaticamente via JS."
            if log_func: log_func(msg)
            else: print(msg)
            dismissed = True
    except Exception:
        pass

    # 3. Modal/Popup do APEX na página via Selenium
    try:
        dialogs = driver.find_elements(By.XPATH, "//div[contains(@class,'ui-dialog') and not(contains(@class,'ui-dialog-popuplov')) and not(contains(@style,'display: none'))]")
        for d in dialogs:
            try:
                txt = d.text.lower()
                if any(k in txt for k in ["beg_vw_transp_pref_id", "identificador", "ora-", "transportadora"]):
                    botoes = d.find_elements(By.XPATH, ".//div[contains(@class,'ui-dialog-buttonpane')]//button | .//button")
                    for b in botoes:
                        if b.is_displayed():
                            driver.execute_script("arguments[0].click();", b)
                            msg = f"   ℹ️ [ERP] Popup de aviso do ERP confirmado automaticamente via Selenium: '{d.text.strip()[:60]}...'"
                            if log_func: log_func(msg)
                            else: print(msg)
                            dismissed = True
                            break
            except Exception:
                pass
    except Exception:
        pass

    return dismissed

def wait_for_apex_ready(driver, timeout=10, log_func=None):
    """Aguarda o Oracle APEX e jQuery terminarem todas as requisições AJAX e animações."""
    dismiss_apex_alert_if_any(driver, log_func=log_func)
    t_end = time.time() + timeout
    time.sleep(0.2)
    while time.time() < t_end:
        dismiss_apex_alert_if_any(driver, log_func=log_func)
        try:
            is_ajax_active = driver.execute_script("return (typeof jQuery !== 'undefined') ? (jQuery.active > 0) : false;")
            has_spinner = driver.execute_script("""
                return document.querySelectorAll('.u-Processing, .apex_wait_overlay, .a-LoadingIndicator').length > 0;
            """)
            if not is_ajax_active and not has_spinner:
                break
        except Exception:
            pass
        time.sleep(0.15)
    time.sleep(0.15)
    dismiss_apex_alert_if_any(driver, log_func=log_func)

def robust_wait_for(wait, by, locator, condition=EC.presence_of_element_located):
    try:
        return wait.until(condition((by, locator)))
    except TimeoutException:
        raise Exception(f"Elemento não encontrado: {locator}")

def robust_click(driver, by, locator, wait_time=20):
    wait_for_apex_ready(driver, timeout=5)
    wait = WebDriverWait(driver, wait_time)
    try:
        el = robust_wait_for(wait, by, locator, EC.presence_of_element_located)
        driver.execute_script("arguments[0].scrollIntoView({behavior: 'auto', block: 'center', inline: 'center'});", el)
        time.sleep(0.1) 
        el = robust_wait_for(wait, by, locator, EC.element_to_be_clickable)
        el.click()
    except Exception:
        try:
            el = driver.find_element(by, locator) 
            driver.execute_script("arguments[0].scrollIntoView({behavior: 'auto', block: 'center', inline: 'center'});", el)
            time.sleep(0.1)
            driver.execute_script("arguments[0].click();", el)
        except Exception as e:
            raise Exception(f"Não consegui clicar em {locator}. Erro: {e}")

def robust_find_and_send_keys(driver, by, locator, keys, wait_time=20):
    wait_for_apex_ready(driver, timeout=5)
    wait = WebDriverWait(driver, wait_time)
    try:
        el = robust_wait_for(wait, by, locator, EC.visibility_of_element_located)
        driver.execute_script("arguments[0].scrollIntoView({behavior: 'auto', block: 'center', inline: 'center'});", el)
        time.sleep(0.1)
        el.clear()
        time.sleep(0.05)
        el.send_keys(str(keys))
    except Exception as e:
        raise Exception(f"Erro ao digitar em {locator}: {e}")

def robust_switch_to_iframe(driver, iframe_tuple, check_tuple, wait_time=20):
    wait = WebDriverWait(driver, wait_time)
    try:
        driver.switch_to.default_content()
        time.sleep(0.2)
        robust_wait_for(wait, iframe_tuple[0], iframe_tuple[1], EC.frame_to_be_available_and_switch_to_it)
        robust_wait_for(wait, check_tuple[0], check_tuple[1], EC.presence_of_element_located)
        time.sleep(0.15)
    except Exception as e:
        raise Exception(f"Erro ao entrar no iframe: {e}")

def robust_check_element(driver, by, locator, wait_time=15, log_func=print):
    wait = WebDriverWait(driver, wait_time)
    try:
        el = robust_wait_for(wait, by, locator, EC.presence_of_element_located)
        driver.execute_script("arguments[0].scrollIntoView({behavior: 'auto', block: 'center', inline: 'center'});", el)
        time.sleep(0.05)
        if "label" in locator:
            el.click()
        elif not el.is_selected():
            try: el.click()
            except: driver.execute_script("arguments[0].click();", el)
        if log_func: log_func(f"   > Elemento marcado: {locator}")
    except Exception as e:
        if log_func: log_func(f"   ! Erro ao marcar elemento {locator}: {e}")

# --- CLASSE PRINCIPAL ---

class AutomacaoSchulzApp:
    def __init__(self, root_window):
        self.root = root_window
        self.root.title("Automação Schulz Unificada")
        self.root.geometry("560x560")
        self.root.attributes('-topmost', True)
        self.driver_lf = None
        self.pedidos_processados = [] 
        self.itens_nao_localizados = []
        self.config = carregar_config()
        self.tabela_path = self.config.get("tabela_schulz_path", "")
        self.criar_tela_inicial()

    def obter_texto_tabela(self):
        if self.tabela_path and os.path.exists(self.tabela_path):
            nome = os.path.basename(self.tabela_path)
            return nome if len(nome) <= 42 else nome[:39] + "..."
        return "Nenhuma tabela configurada (clique ao lado)"

    def restaurar_botoes_iniciais(self):
        self.btn_iniciar.config(state=tk.NORMAL, text="Carregar Pedidos e Iniciar")
        self.btn_alterar_tabela.config(state=tk.NORMAL)
        self.radio_auto.config(state=tk.NORMAL)
        self.radio_comp.config(state=tk.NORMAL)
        self.rb_normal.config(state=tk.NORMAL)
        self.rb_teste.config(state=tk.NORMAL)
        self.entry_inicio.config(state=tk.NORMAL)
        self.entry_coluna.config(state=tk.NORMAL)
        self.root.update()

    def alterar_tabela_schulz(self):
        novo_path = filedialog.askopenfilename(
            title="Selecione a Tabela Schulz (Preços e Códigos)", 
            filetypes=[("Excel", "*.xlsx *.xls")]
        )
        if novo_path:
            try:
                df_teste = pd.read_excel(novo_path, header=None, nrows=10)
                txt_teste = ' '.join([str(x).lower() for x in df_teste.values.flatten() if pd.notna(x)])
                if ('o.c.' in txt_teste or 'saldo ordem' in txt_teste) and 'feedback' in novo_path.lower():
                    self.force_top_window()
                    messagebox.showwarning(
                        "Aviso de Arquivo",
                        f"⚠️ Atenção:\nO arquivo selecionado parece ser uma Planilha de Pedidos (FeedBack):\n'{os.path.basename(novo_path)}'\n\n"
                        f"Neste campo você deve selecionar a Tabela Schulz com os códigos LF e preços negociados.",
                        parent=self.root
                    )
            except Exception:
                pass

            self.tabela_path = novo_path
            self.config["tabela_schulz_path"] = novo_path
            salvar_config(self.config)
            self.lbl_tabela_nome.config(text=self.obter_texto_tabela())
            self.log(f"📁 Tabela Schulz configurada: {os.path.basename(novo_path)}")

    def criar_tela_inicial(self):
        # Frame de Configurações
        self.frame_config = tk.Frame(self.root)
        self.frame_config.pack(pady=10, fill=tk.X)

        tk.Label(self.frame_config, text="Automação Schulz Unificada", font=("Arial", 12, "bold")).pack(pady=5)
        
        self.empresa_var = tk.StringVar(value="Automotiva")
        frame_radios = tk.Frame(self.frame_config)
        frame_radios.pack(pady=5)
        self.radio_auto = tk.Radiobutton(frame_radios, text="Schulz Automotiva", variable=self.empresa_var, value="Automotiva")
        self.radio_auto.pack(side=tk.LEFT, padx=10)
        self.radio_comp = tk.Radiobutton(frame_radios, text="Schulz Compressores", variable=self.empresa_var, value="Compressores")
        self.radio_comp.pack(side=tk.LEFT, padx=10)

        # Frame Configuração Tabela Schulz
        frame_tabela = tk.LabelFrame(self.frame_config, text=" Tabela Schulz (Preços e Códigos) ", font=("Arial", 9, "bold"), padx=10, pady=6)
        frame_tabela.pack(pady=4, fill=tk.X, padx=25)

        self.lbl_tabela_nome = tk.Label(frame_tabela, text=self.obter_texto_tabela(), fg="#1565C0", font=("Arial", 8, "italic"), anchor="w")
        self.lbl_tabela_nome.pack(side=tk.LEFT, fill=tk.X, expand=True, padx=(0, 5))

        self.btn_alterar_tabela = tk.Button(frame_tabela, text="Alterar Planilha...", command=self.alterar_tabela_schulz, font=("Arial", 8, "bold"), bg="#e0e0e0")
        self.btn_alterar_tabela.pack(side=tk.RIGHT)

        # Frame Opções de Execução (Modo de Execução, Pedido Inicial e Coluna)
        frame_modos = tk.LabelFrame(self.frame_config, text=" Opções de Execução ", font=("Arial", 9, "bold"), padx=10, pady=5)
        frame_modos.pack(pady=4, fill=tk.X, padx=25)

        # Radios Modo Normal / Teste
        self.modo_exec_var = tk.StringVar(value="Normal")
        frame_radios_modo = tk.Frame(frame_modos)
        frame_radios_modo.pack(fill=tk.X, pady=2)
        self.rb_normal = tk.Radiobutton(frame_radios_modo, text="Modo Normal (Todos)", variable=self.modo_exec_var, value="Normal", font=("Arial", 9, "bold"))
        self.rb_normal.pack(side=tk.LEFT, padx=5)
        self.rb_teste = tk.Radiobutton(frame_radios_modo, text="Modo Teste (Apenas 1 pedido)", variable=self.modo_exec_var, value="Teste", fg="#D32F2F", font=("Arial", 9, "bold"))
        self.rb_teste.pack(side=tk.LEFT, padx=15)

        # Linha de Pedido Inicial e Coluna de Preço
        frame_opcoes_linha = tk.Frame(frame_modos)
        frame_opcoes_linha.pack(fill=tk.X, pady=3)

        tk.Label(frame_opcoes_linha, text="Começar do Pedido nº:").pack(side=tk.LEFT, padx=2)
        self.inicio_pedido_var = tk.StringVar(value="1")
        self.entry_inicio = tk.Entry(frame_opcoes_linha, textvariable=self.inicio_pedido_var, width=4, justify="center")
        self.entry_inicio.pack(side=tk.LEFT, padx=3)

        tk.Label(frame_opcoes_linha, text="Coluna Preço:").pack(side=tk.LEFT, padx=(15, 2))
        self.coluna_preco_var = tk.StringVar(value=self.config.get("coluna_preco", "I"))
        self.entry_coluna = tk.Entry(frame_opcoes_linha, textvariable=self.coluna_preco_var, width=4, justify="center")
        self.entry_coluna.pack(side=tk.LEFT, padx=3)

        self.btn_iniciar = tk.Button(self.frame_config, text="Carregar Pedidos e Iniciar", command=self.iniciar_processo, height=2, bg="#4CAF50", fg="white", font=("Arial", 10, "bold"))
        self.btn_iniciar.pack(pady=8, fill=tk.X, padx=30)

        # Botão para fechar o navegador
        self.btn_fechar_chrome = tk.Button(self.frame_config, text="Fechar Navegador Chrome", command=self.fechar_driver, state=tk.DISABLED, bg="#f44336", fg="white")
        self.btn_fechar_chrome.pack(pady=2, fill=tk.X, padx=30)

        # Frame de Logs
        self.frame_logs = tk.Frame(self.root)
        self.frame_logs.pack(pady=5, fill=tk.BOTH, expand=True, padx=15)
        
        tk.Label(self.frame_logs, text="Painel de Execução / Logs:", font=("Arial", 9, "bold")).pack(anchor=tk.W, pady=2)
        
        self.log_text = tk.Text(self.frame_logs, height=12, state=tk.DISABLED, wrap=tk.WORD, bg="#1e1e1e", fg="#d4d4d4", font=("Consolas", 9))
        self.log_text.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)

        scrollbar = tk.Scrollbar(self.frame_logs, command=self.log_text.yview)
        scrollbar.pack(side=tk.RIGHT, fill=tk.Y)
        self.log_text.config(yscrollcommand=scrollbar.set)

    def log(self, message):
        print(message)
        self.log_text.config(state=tk.NORMAL)
        self.log_text.insert(tk.END, f"{message}\n")
        self.log_text.see(tk.END)
        self.log_text.config(state=tk.DISABLED)
        self.root.update_idletasks()
        self.root.update()

    def fechar_driver(self):
        if self.driver_lf:
            try:
                self.driver_lf.quit()
                self.log("ℹ️ Navegador Chrome encerrado com sucesso.")
            except Exception as e:
                self.log(f"⚠️ Erro ao fechar o navegador: {e}")
            self.driver_lf = None
            self.btn_fechar_chrome.config(state=tk.DISABLED)
            self.root.update()

    def force_top_window(self):
        self.root.deiconify()
        self.root.attributes('-topmost', True)
        self.root.lift()

    def input_topmost(self, title, prompt):
        self.force_top_window()
        return simpledialog.askstring(title, prompt, parent=self.root)

    def iniciar_processo(self):
        path_pedidos = filedialog.askopenfilename(title="Selecione a Planilha de PEDIDOS", filetypes=[("Excel", "*.xlsx *.xls")])
        if not path_pedidos: return

        if not self.tabela_path or not os.path.exists(self.tabela_path):
            self.force_top_window()
            messagebox.showinfo("Configurar Tabela", "Por favor, selecione a Tabela Schulz (com Preços e Códigos).")
            self.alterar_tabela_schulz()
            if not self.tabela_path or not os.path.exists(self.tabela_path):
                return
        
        path_codigos = self.tabela_path

        # Desabilitar botões durante execução
        self.btn_iniciar.config(state=tk.DISABLED, text="Executando Automação...")
        self.btn_alterar_tabela.config(state=tk.DISABLED)
        self.radio_auto.config(state=tk.DISABLED)
        self.radio_comp.config(state=tk.DISABLED)
        self.rb_normal.config(state=tk.DISABLED)
        self.rb_teste.config(state=tk.DISABLED)
        self.entry_inicio.config(state=tk.DISABLED)
        self.entry_coluna.config(state=tk.DISABLED)
        self.btn_fechar_chrome.config(state=tk.DISABLED)
        self.root.update()

        try:
            col_letra = self.coluna_preco_var.get().strip().upper()
            if not col_letra.isalpha():
                col_letra = "I"
            self.config["coluna_preco"] = col_letra
            salvar_config(self.config)

            self.log(f"📄 Tabela Schulz utilizada: {os.path.basename(path_codigos)} (Coluna {col_letra})")

            # --- 1. VALIDAÇÃO E LEITURA INTELIGENTE DA PLANILHA DE PEDIDOS ---
            df_ped_raw = pd.read_excel(path_pedidos, header=None)
            
            # Detectar linha de cabeçalho
            header_ped_idx = None
            for r in range(min(10, len(df_ped_raw))):
                row_vals = [str(x).strip().lower() for x in df_ped_raw.iloc[r] if pd.notna(x)]
                if any(w in ['o.c.', 'oc', 'o.c'] or 'ordem' in w for w in row_vals):
                    header_ped_idx = r
                    break
            
            # Verificação se o usuário selecionou uma tabela de preços por engano
            todas_palavras_ped = ' '.join([str(x).lower() for x in df_ped_raw.iloc[:10].values.flatten() if pd.notna(x)])
            if header_ped_idx is None and any(w in todas_palavras_ped for w in ['cod lf', 'categoria', 'vencido', 'tabela lf', 'contrato']):
                raise ValueError(
                    f"O arquivo selecionado parece ser uma TABELA DE PREÇOS:\n'{os.path.basename(path_pedidos)}'\n\n"
                    f"⚠️ Por favor, selecione a PLANILHA DE PEDIDOS (FeedBack) da Schulz!\n"
                    f"(Exemplo: 'Schulz Automotiva ... FeedBack.xlsx')."
                )
            
            if header_ped_idx is None:
                header_ped_idx = 0
                
            headers_ped = [str(x).strip().lower() if pd.notna(x) else '' for x in df_ped_raw.iloc[header_ped_idx]]
            
            # Mapeamento dinâmico de colunas
            col_item_idx = None
            col_oc_idx = None
            col_qtd_idx = None
            col_estab_idx = None
            
            for idx, h in enumerate(headers_ped):
                if h == 'item' or (col_item_idx is None and ('item' in h or 'cod' in h) and 'descri' not in h):
                    col_item_idx = idx
                elif h in ['o.c.', 'oc', 'o.c'] or (col_oc_idx is None and ('ordem de compra' in h or 'pedido' in h)):
                    col_oc_idx = idx
                elif h in ['qtd ordem', 'quantidade'] or (col_qtd_idx is None and ('qtd' in h or 'saldo' in h)):
                    col_qtd_idx = idx
                elif 'estab' in h or 'filial' in h:
                    col_estab_idx = idx
            
            # Fallbacks tradicionais para colunas padrão (C=2, F=5, K=10, O=14)
            if col_item_idx is None: col_item_idx = 2
            if col_oc_idx is None:   col_oc_idx = 5
            if col_qtd_idx is None:  col_qtd_idx = 10
            if col_estab_idx is None and df_ped_raw.shape[1] > 14: col_estab_idx = 14
            
            letra_item = get_column_letter(col_item_idx + 1)
            letra_oc = get_column_letter(col_oc_idx + 1)
            letra_qtd = get_column_letter(col_qtd_idx + 1)
            letra_estab = get_column_letter(col_estab_idx + 1) if col_estab_idx is not None else "N/A"
            
            cols_to_use = [col_item_idx, col_oc_idx, col_qtd_idx]
            col_names = ['cod_item', 'ordem_compra', 'qtd']
            if col_estab_idx is not None and col_estab_idx < df_ped_raw.shape[1]:
                cols_to_use.append(col_estab_idx)
                col_names.append('estab_ent')
            
            df_ped = df_ped_raw.iloc[header_ped_idx + 1:, cols_to_use].copy()
            df_ped.columns = col_names
            if 'estab_ent' not in df_ped.columns:
                df_ped['estab_ent'] = "1 - SCHULZ S.A."
                
            df_ped = df_ped.dropna(subset=['ordem_compra', 'cod_item']).reset_index(drop=True)
            if df_ped.empty:
                raise ValueError("Nenhum pedido válido encontrado na planilha selecionada.")

            # Validação se a coluna de OC não contém preços
            amostras_oc = [str(x).strip() for x in df_ped['ordem_compra'].dropna().tolist()[:10]]
            for val in amostras_oc:
                if 'r$' in val.lower() or ',' in val:
                    raise ValueError(
                        f"A coluna detectada como Ordem de Compra ({letra_oc}) contém valores no formato monetário: '{val}'.\n"
                        f"Verifique se o arquivo selecionado é realmente o FeedBack de pedidos."
                    )

            # --- 2. VALIDAÇÃO E LEITURA DA TABELA SCHULZ ---
            df_cod_raw = pd.read_excel(path_codigos, header=None)
            num_cols_cod = df_cod_raw.shape[1]
            idx_preco = column_index_from_string(col_letra) - 1
            if idx_preco >= num_cols_cod:
                max_col_letra = get_column_letter(num_cols_cod)
                raise ValueError(
                    f"A coluna de preço configurada '{col_letra}' não existe na Tabela Schulz selecionada.\n"
                    f"A tabela possui colunas de A até {max_col_letra}."
                )

            df_item  = df_cod_raw.iloc[:, 0].to_frame(name='cod_item')
            df_tipo  = df_cod_raw.iloc[:, 2].to_frame(name='tipo_uso') if num_cols_cod > 2 else pd.DataFrame({'tipo_uso': ['CONSUMO'] * len(df_item)})
            df_lf    = df_cod_raw.iloc[:, 3].to_frame(name='cod_lf') if num_cols_cod > 3 else pd.DataFrame({'cod_lf': [''] * len(df_item)})
            df_preco = df_cod_raw.iloc[:, idx_preco].to_frame(name='preco')
            
            df_cod = pd.concat([df_item, df_tipo, df_lf, df_preco], axis=1)
            df_cod = df_cod.iloc[1:].reset_index(drop=True)

            df_ped = df_ped.dropna(subset=['ordem_compra'])
            df_cod = df_cod.dropna(subset=['cod_item', 'cod_lf'])
            df_cod = df_cod.drop_duplicates(subset=['cod_item'], keep='first')
            
            def padronizar_codigo(codigo):
                c = str(codigo).strip()
                if c.endswith('.0'): c = c[:-2]
                if '/' in c: c = c.split('/')[0]
                if len(c) < 6 and c.isdigit(): c = c.zfill(6)
                return c

            df_ped['cod_item'] = df_ped['cod_item'].apply(padronizar_codigo)
            df_ped['ordem_compra'] = df_ped['ordem_compra'].astype(str).str.strip().replace(r'\.0$', '', regex=True)
            if 'estab_ent' in df_ped.columns:
                df_ped['estab_ent'] = df_ped['estab_ent'].astype(str).str.strip()
            
            df_cod['cod_item'] = df_cod['cod_item'].apply(padronizar_codigo)
            df_cod['cod_lf'] = df_cod['cod_lf'].apply(padronizar_codigo)
            df_cod['tipo_uso'] = df_cod['tipo_uso'].astype(str).str.strip()
            
            df_ped['qtd'] = pd.to_numeric(df_ped['qtd'], errors='coerce').fillna(0)
            
            def limpar_preco(valor):
                if pd.isna(valor): return 0.0
                v_str = str(valor).upper().replace('R$', '').replace(' ', '').strip()
                if ',' in v_str and '.' in v_str:
                    v_str = v_str.replace('.', '').replace(',', '.')
                elif ',' in v_str:
                    v_str = v_str.replace(',', '.')
                try: return float(v_str)
                except: return 0.0
            df_cod['preco'] = df_cod['preco'].apply(limpar_preco)

            # Cruzamento
            self.pedidos_processados = pd.merge(df_ped, df_cod, on='cod_item', how='left')
            
            # --- 3. RESUMO E CONFIRMAÇÃO DE PRÉ-EXECUÇÃO ---
            total_itens = len(self.pedidos_processados)
            ocs_unicas = self.pedidos_processados['ordem_compra'].unique().tolist()
            total_ocs = len(ocs_unicas)
            exemplo_ocs = ", ".join(ocs_unicas[:4]) + ("..." if total_ocs > 4 else "")
            
            itens_sem_lf = self.pedidos_processados[self.pedidos_processados['cod_lf'].isna() | (self.pedidos_processados['cod_lf'] == '')]
            qtd_sem_lf = len(itens_sem_lf)

            self.log("=" * 45)
            self.log("🔍 RESUMO DA VERIFICAÇÃO PRÉ-EXECUÇÃO:")
            self.log(f"  • Planilha de Pedidos: {os.path.basename(path_pedidos)}")
            self.log(f"    - Coluna Item: {letra_item} | Coluna OC: {letra_oc} | Coluna Qtd: {letra_qtd}")
            self.log(f"    - Total de OCs: {total_ocs} ({exemplo_ocs})")
            self.log(f"    - Total de Itens: {total_itens}")
            self.log(f"  • Tabela Schulz: {os.path.basename(path_codigos)}")
            self.log(f"    - Coluna Preço: {col_letra} | Códigos cadastrados: {len(df_cod)}")
            if qtd_sem_lf > 0:
                self.log(f"  ⚠️ Atenção: {qtd_sem_lf} item(ns) estão sem Código LF cadastrado na tabela.")
            self.log("=" * 45)

            # Popup de confirmação antes de abrir o navegador
            self.force_top_window()
            msg_confirmacao = (
                f"📋 CONFERÊNCIA ANTES DE INICIAR:\n\n"
                f"• Planilha de Pedidos: {os.path.basename(path_pedidos)}\n"
                f"  - Ordens de Compra (OCs): {total_ocs} (Ex: {exemplo_ocs})\n"
                f"  - Coluna da OC detectada: Coluna {letra_oc}\n"
                f"  - Total de Itens a incluir: {total_itens}\n\n"
                f"• Tabela Schulz: {os.path.basename(path_codigos)}\n"
                f"  - Coluna de Preço: Coluna {col_letra}\n"
            )
            if qtd_sem_lf > 0:
                msg_confirmacao += f"\n⚠️ Nota: {qtd_sem_lf} item(ns) não possuem Código LF correspondente (serão gravados no log de itens não localizados).\n"
            msg_confirmacao += "\nTudo correto? Deseja iniciar a automação no Portal LF agora?"

            if not messagebox.askyesno("Confirmar Início da Automação", msg_confirmacao, parent=self.root):
                self.log("⏹️ Automação cancelada pelo usuário após a conferência.")
                self.restaurar_botoes_iniciais()
                return

            self.log("📋 Dados conferidos e confirmados! Iniciando navegador...")
            self.iniciar_automacao()
            
        except Exception as e:
            self.force_top_window()
            messagebox.showerror("Erro", f"Erro planilhas: {e}")
            self.restaurar_botoes_iniciais()

    def iniciar_automacao(self):
        try:
            options = webdriver.ChromeOptions()
            options.add_argument("--start-maximized")
            self.driver_lf = webdriver.Chrome(options=options)
            self.btn_fechar_chrome.config(state=tk.NORMAL)
            
            self.log("🔑 Abrindo portal LF...")
            self.driver_lf.get(URL_PORTAL_LF)
            
            wait_login = WebDriverWait(self.driver_lf, 25)
            user_el = wait_login.until(EC.visibility_of_element_located((By.XPATH, XPATHS_LF["login_user"])))
            time.sleep(0.8) # Pausa para inicialização dos scripts da página
            
            self.log("🔑 Preenchendo usuário e senha...")
            user_el.clear()
            user_el.send_keys(CREDENCIAIS["LF"]["user"])
            time.sleep(0.4)
            
            pass_el = wait_login.until(EC.visibility_of_element_located((By.XPATH, XPATHS_LF["login_pass"])))
            pass_el.clear()
            pass_el.send_keys(CREDENCIAIS["LF"]["pass"])
            time.sleep(0.4)
            
            # Dispara blur/tab para o APEX processar e validar o campo de senha
            pass_el.send_keys(Keys.TAB)
            time.sleep(0.8)
            
            self.log("🔑 Confirmando login...")
            btn_login = wait_login.until(EC.element_to_be_clickable((By.XPATH, XPATHS_LF["login_btn"])))
            try:
                btn_login.click()
            except Exception:
                self.driver_lf.execute_script("arguments[0].click();", btn_login)
            
            # Aguarda o menu principal aparecer
            try:
                WebDriverWait(self.driver_lf, 35).until(EC.presence_of_element_located((By.XPATH, XPATHS_LF["nav_menu_btn"])))
                self.log("✅ Login efetuado com sucesso!")
            except TimeoutException:
                # Se ainda estiver na tela de login, tenta mais um clique de segurança
                botoes = self.driver_lf.find_elements(By.XPATH, XPATHS_LF["login_btn"])
                if botoes and botoes[0].is_displayed():
                    self.log("⚠️ Reenviando confirmação de login...")
                    try: botoes[0].click()
                    except: self.driver_lf.execute_script("arguments[0].click();", botoes[0])
                    WebDriverWait(self.driver_lf, 25).until(EC.presence_of_element_located((By.XPATH, XPATHS_LF["nav_menu_btn"])))
                    self.log("✅ Login efetuado com sucesso!")
                else:
                    raise Exception("Tempo esgotado aguardando o menu principal após login.")
            
            self.loop_pedidos()
            
        except Exception as e:
            self.log(f"❌ Erro na automação: {e}")
            messagebox.showerror("Erro Login", f"{e}")
        finally:
            # Reabilitar botões após conclusão ou erro
            self.btn_iniciar.config(state=tk.NORMAL, text="Carregar Pedidos e Iniciar")
            self.btn_alterar_tabela.config(state=tk.NORMAL)
            self.radio_auto.config(state=tk.NORMAL)
            self.radio_comp.config(state=tk.NORMAL)
            self.rb_normal.config(state=tk.NORMAL)
            self.rb_teste.config(state=tk.NORMAL)
            self.entry_inicio.config(state=tk.NORMAL)
            self.entry_coluna.config(state=tk.NORMAL)
            self.root.update()

    def loop_pedidos(self):
        grupos = list(self.pedidos_processados.groupby('ordem_compra', sort=False))
        total_grupos = len(grupos)
        self.log(f"📦 Encontrados {total_grupos} pedido(s) para processar.")

        # Validar número inicial
        try:
            inicio_num = int(self.inicio_pedido_var.get().strip())
            if inicio_num < 1:
                inicio_num = 1
        except Exception:
            inicio_num = 1

        if inicio_num > total_grupos:
            self.log(f"⚠️ Pedido inicial ({inicio_num}) é maior que o total ({total_grupos}). Iniciando do pedido 1.")
            inicio_num = 1

        pedidos_para_processar = grupos[inicio_num - 1:]

        eh_modo_teste = (self.modo_exec_var.get() == "Teste")
        if eh_modo_teste:
            pedidos_para_processar = pedidos_para_processar[:1]
            self.log(f"🧪 [MODO TESTE] Processando APENAS 1 pedido (Pedido nº {inicio_num} de {total_grupos}).")
        else:
            self.log(f"▶️ Processando {len(pedidos_para_processar)} pedido(s), iniciando no pedido nº {inicio_num} de {total_grupos}.")

        for i, (ordem_compra, itens) in enumerate(pedidos_para_processar, 1):
            idx_real = inicio_num + i - 1
            try:
                self.log(f"\n🚀 [{idx_real}/{total_grupos}] Processando Pedido: {ordem_compra}")
                
                tipo_uso_pedido = itens.iloc[0]['tipo_uso']
                valor_total_pedido = (itens['qtd'] * itens['preco']).sum()
                
                estab_raw = str(itens.iloc[0].get('estab_ent', '1 - SCHULZ S.A.')).strip()
                estab_upper = estab_raw.upper()
                eh_joinville = (
                    estab_upper in ["1", "1 - SCHULZ S.A.", "SCHULZ S.A."]
                    or "JOINVILLE" in estab_upper
                    or estab_upper.startswith("1 -")
                    or estab_upper.startswith("1 ")
                    or estab_upper == ""
                    or estab_upper == "NAN"
                )
                
                str_tipo = str(tipo_uso_pedido).upper()
                eh_industrializacao = ("CONSUMO" not in str_tipo) if ("NAN" not in str_tipo and str_tipo.strip() != "") else True

                self.log(f"   > Tipo: {tipo_uso_pedido} | Estab: {estab_raw} (Joinville={eh_joinville}) | Valor: R$ {valor_total_pedido:.2f}")

                self.driver_lf.switch_to.default_content()
                try: 
                    robust_click(self.driver_lf, By.XPATH, XPATHS_LF["nav_pedido_link"], wait_time=3)
                except:
                    robust_click(self.driver_lf, By.XPATH, XPATHS_LF["nav_menu_btn"])
                    time.sleep(0.3)
                    robust_click(self.driver_lf, By.XPATH, XPATHS_LF["nav_pedido_link"])
                
                # Aguardar o carregamento da lista de pedidos e tentar clicar em "Novo Pedido"
                btn_clicado = False
                for xpath in XPATHS_LF["novo_pedido_btn_opcoes"]:
                    try:
                        robust_click(self.driver_lf, By.XPATH, xpath, wait_time=3)
                        btn_clicado = True
                        break
                    except: 
                        continue
                if not btn_clicado: 
                    raise Exception("Botão Novo não encontrado.")

                # Preencher Cabeçalho
                cnpj_selecionado = CNPJS_SCHULZ[self.empresa_var.get()]
                self.preencher_cabecalho(ordem_compra, tipo_uso_pedido, valor_total_pedido, cnpj_selecionado, eh_joinville=eh_joinville)
                
                # Itens
                self.preencher_itens(itens, ordem_compra)
                
                self.log(f"✅ Pedido {ordem_compra} concluído com sucesso!")

            except Exception as e:
                self.log(f"❌ Erro no pedido {ordem_compra}: {e}")
                if self.driver_lf:
                    try: 
                        self.driver_lf.find_element(By.CSS_SELECTOR, ".ui-dialog-titlebar-close").click()
                    except: 
                        pass
                    try:
                        self.driver_lf.switch_to.default_content()
                    except:
                        pass
                
                self.force_top_window()
                if not messagebox.askyesno("Erro", f"Erro no pedido {ordem_compra}:\n{e}\n\nDeseja continuar com o próximo pedido?"): 
                    break
        
        if eh_modo_teste:
            self.log("\n🧪 [MODO TESTE] Finalizado! Pedido de teste processado com sucesso.")
            self.force_top_window()
            messagebox.showinfo("Modo Teste Concluído", "Pedido de teste processado com sucesso!\n\nVerifique no portal LF se tudo foi preenchido corretamente.")
        elif self.itens_nao_localizados:
            relatorio = "\n".join(self.itens_nao_localizados)
            with open("itens_nao_localizados.txt", "w", encoding="utf-8") as f:
                f.write("Itens não localizados ou com erro:\n\n")
                f.write(relatorio)
            self.log("\n⚠️ Algumas linhas de itens não puderam ser processadas. Veja itens_nao_localizados.txt")
            self.force_top_window()
            messagebox.showwarning("Aviso", f"Concluído com avisos.\nAlguns itens não foram localizados (gerado 'itens_nao_localizados.txt').\n\n{relatorio[:300]}...")
        else:
            self.log("\n🎉 Processo concluído! Todos os pedidos foram enviados perfeitamente.")
            self.force_top_window()
            messagebox.showinfo("Fim", "Concluído com sucesso!")

    # --- FUNÇÃO DE POPUP ---
    def abrir_popup_e_selecionar(self, valor_busca, exact_match=False, ignorar_erro=False, btn_reabrir=None, iframe_info=None):
        wait = WebDriverWait(self.driver_lf, 20)
        xpath_popup = "//div[contains(@class,'ui-dialog-popuplov') and not(contains(@style,'display: none'))]"
        
        # Garante que o popup abriu. Se não abriu (devido a AJAX do APEX ter engolido o clique), tenta reabrir até 3x
        popup_aberto = False
        for tentativa in range(3):
            dismiss_apex_alert_if_any(self.driver_lf, log_func=self.log)
            self.driver_lf.switch_to.default_content()
            try:
                WebDriverWait(self.driver_lf, 2.5).until(EC.visibility_of_element_located((By.XPATH, xpath_popup)))
                popup_aberto = True
                break
            except TimeoutException:
                dismiss_apex_alert_if_any(self.driver_lf, log_func=self.log)
                if btn_reabrir:
                    self.log(f"   [POPUP] ⚠️ Popup não detectado (tentativa {tentativa+1}/3). Reabrindo...")
                    if iframe_info:
                        robust_switch_to_iframe(self.driver_lf, iframe_info[0], iframe_info[1])
                    else:
                        robust_switch_to_iframe(self.driver_lf, 
                                                (By.XPATH, XPATHS_LF["novo_pedido_iframe"]), 
                                                (By.ID, XPATHS_LF["novo_pedido_iframe_check"]))
                    wait_for_apex_ready(self.driver_lf, timeout=5, log_func=self.log)
                    robust_click(self.driver_lf, btn_reabrir[0], btn_reabrir[1])
                    time.sleep(0.4)
                else:
                    time.sleep(0.5)

        if not popup_aberto:
            # Tenta uma última espera com timeout maior caso a rede esteja lenta
            self.driver_lf.switch_to.default_content()
            try:
                WebDriverWait(self.driver_lf, 4).until(EC.visibility_of_element_located((By.XPATH, xpath_popup)))
                popup_aberto = True
            except TimeoutException:
                pass

        if not popup_aberto:
            if ignorar_erro:
                return False
            self.log(f"   [POPUP] ⚠️ O popup de pesquisa para '{valor_busca}' não abriu.")
            self.force_top_window()
            if messagebox.askyesno("Ajuda Manual", f"O popup para pesquisar '{valor_busca}' não abriu automaticamente.\n\nVocê já abriu ou deseja abrir manualmente no navegador e prosseguir?\nClique 'SIM' para continuar, ou 'NÃO' para abortar este pedido."):
                self.driver_lf.switch_to.default_content()
                try:
                    WebDriverWait(self.driver_lf, 5).until(EC.visibility_of_element_located((By.XPATH, xpath_popup)))
                    popup_aberto = True
                except Exception:
                    return True
            else:
                raise Exception(f"Popup para '{valor_busca}' não abriu no portal.")

        xpath_input = f"{xpath_popup}//input[contains(@class, 'a-PopupLOV-search')]"
        xpath_btn_search = f"{xpath_popup}//button[contains(@class,'a-PopupLOV-doSearch')]"
        xpath_load_more = f"{xpath_popup}//button[contains(text(), 'Load More') or contains(text(), 'Carregar Mais')]"
        
        inp = wait.until(EC.visibility_of_element_located((By.XPATH, xpath_input)))
        try:
            inp.click()
            inp.send_keys(Keys.CONTROL + "a")
            inp.send_keys(Keys.BACKSPACE)
            time.sleep(0.1)
        except:
            pass
        inp.clear()
        
        valor_limpo = str(valor_busca).strip()
        if valor_limpo.endswith('.0'):
            valor_limpo = valor_limpo[:-2]
        self.log(f"   [POPUP] Pesquisando por: '{valor_limpo}' (Exato={exact_match})")
        
        inp.send_keys(valor_limpo)
        time.sleep(0.2)
        inp.send_keys(Keys.RETURN)
        try: 
            btn_s = self.driver_lf.find_elements(By.XPATH, xpath_btn_search)
            if btn_s and btn_s[0].is_displayed():
                btn_s[0].click()
        except: 
            pass

        # Pausa inicial para o Oracle APEX carregar os resultados da busca via AJAX
        time.sleep(1.2)

        tempo_fim = time.time() + 10.0
        encontrado = False

        cod_sem_zeros = valor_limpo.lstrip('0') if valor_limpo.isdigit() else valor_limpo

        while time.time() < tempo_fim and not encontrado:
            # 1. Procura primeiro nas linhas da tabela de resultados
            try:
                linhas = self.driver_lf.find_elements(By.XPATH, f"{xpath_popup}//div[contains(@class, 'a-PopupLOV-results')]//table//tbody//tr")
                for tr in linhas:
                    if tr.is_displayed():
                        txt = tr.text.strip()
                        if (valor_limpo in txt) or (cod_sem_zeros and cod_sem_zeros in txt):
                            self.driver_lf.execute_script("arguments[0].scrollIntoView({behavior: 'auto', block: 'center', inline: 'center'});", tr)
                            time.sleep(0.1)
                            tds = tr.find_elements(By.TAG_NAME, "td")
                            alvo_click = tds[0] if tds else tr
                            try:
                                alvo_click.click()
                            except:
                                self.driver_lf.execute_script("arguments[0].click();", alvo_click)
                            self.log(f"   [POPUP] ✅ Selecionado com sucesso via linha ({valor_limpo}).")
                            time.sleep(0.4)
                            try:
                                WebDriverWait(self.driver_lf, 4).until(EC.invisibility_of_element_located((By.XPATH, xpath_popup)))
                            except:
                                pass
                            return True
            except Exception:
                pass

            # 2. XPaths específicos por célula
            if exact_match:
                xpath_alvo = [
                    f"{xpath_popup}//td[normalize-space(.)='{valor_limpo}']",
                    f"{xpath_popup}//li[normalize-space(.)='{valor_limpo}']",
                    f"{xpath_popup}//a[normalize-space(.)='{valor_limpo}']"
                ]
            else:
                xpath_alvo = [
                    f"{xpath_popup}//td[normalize-space(.)='{valor_limpo}']",
                    f"{xpath_popup}//a[normalize-space(.)='{valor_limpo}']",
                    f"{xpath_popup}//td[normalize-space(.)='{cod_sem_zeros}']",
                    f"{xpath_popup}//a[normalize-space(.)='{cod_sem_zeros}']",
                    f"{xpath_popup}//td[contains(normalize-space(.), '{valor_limpo}')]",
                    f"{xpath_popup}//a[contains(normalize-space(.), '{valor_limpo}')]",
                    f"{xpath_popup}//td[contains(normalize-space(.), '{cod_sem_zeros}')]",
                    f"{xpath_popup}//a[contains(normalize-space(.), '{cod_sem_zeros}')]"
                ]
                # Se já passou pelo menos 2.5 segundos do loop, aceita a primeira linha visível da tabela filtrada
                if time.time() > (tempo_fim - 7.5):
                    xpath_alvo.extend([
                        f"{xpath_popup}//div[contains(@class, 'a-PopupLOV-results')]//table//tbody//tr[1]//td[1]",
                        f"{xpath_popup}//div[contains(@class, 'a-PopupLOV-results')]//table//tr[1]//td[1]",
                        f"{xpath_popup}//div[contains(@class, 'a-PopupLOV-results')]//li[1]"
                    ])

            for xpath in xpath_alvo:
                try:
                    elementos = self.driver_lf.find_elements(By.XPATH, xpath)
                    for el in elementos:
                        if el.is_displayed():
                            self.driver_lf.execute_script("arguments[0].scrollIntoView({behavior: 'auto', block: 'center', inline: 'center'});", el)
                            time.sleep(0.1)
                            try: 
                                el.click()
                            except: 
                                self.driver_lf.execute_script("arguments[0].click();", el)
                            
                            self.log(f"   [POPUP] ✅ Selecionado com sucesso.")
                            time.sleep(0.4)
                            try:
                                WebDriverWait(self.driver_lf, 4).until(EC.invisibility_of_element_located((By.XPATH, xpath_popup)))
                            except:
                                pass
                            return True
                except Exception:
                    continue

            # Se ainda não encontrou, tenta clicar no Load More se houver
            try:
                load_more_els = self.driver_lf.find_elements(By.XPATH, xpath_load_more)
                if load_more_els and load_more_els[0].is_displayed():
                    btn_load = load_more_els[0]
                    self.log("   [POPUP] ⚠️ Clicando em 'Load More'...")
                    self.driver_lf.execute_script("arguments[0].scrollIntoView({behavior: 'auto', block: 'center', inline: 'center'});", btn_load)
                    btn_load.click()
                    time.sleep(1.0)
            except Exception:
                pass

            time.sleep(0.4)

        if not encontrado:
            self.log(f"   [POPUP] ⚠️ Não consegui clicar automaticamente em '{valor_limpo}'.")
            self.force_top_window()
            if messagebox.askyesno("Ajuda Manual", f"O item com código '{valor_limpo}' apareceu no portal?\n\nClique em 'SIM' para você clicar nele manualmente no navegador e prosseguir com o pedido.\nOu clique em 'NÃO' para pular este item."):
                messagebox.showinfo("Aguardando", f"Por favor, clique no item '{valor_limpo}' no portal LF e em seguida clique em OK aqui.")
                return True
            return False

    def preencher_cabecalho(self, ordem_compra, tipo_uso, valor_total, cnpj_cliente, eh_joinville=True):
        wait = WebDriverWait(self.driver_lf, 25)
        
        robust_switch_to_iframe(self.driver_lf, 
                                (By.XPATH, XPATHS_LF["novo_pedido_iframe"]),
                                (By.ID, XPATHS_LF["novo_pedido_iframe_check"]))
        wait_for_apex_ready(self.driver_lf)
        
        # 1. FILIAL 22
        try:
            el_filial = robust_wait_for(wait, By.ID, XPATHS_LF["campo_filial"])
            sel_filial = Select(el_filial)
            if sel_filial.first_selected_option.get_attribute("value") != "22":
                sel_filial.select_by_value("22")
                wait_for_apex_ready(self.driver_lf)
        except: pass

        # 2. CLIENTE
        robust_click(self.driver_lf, By.XPATH, XPATHS_LF["btn_cliente_lov"])
        self.abrir_popup_e_selecionar(cnpj_cliente, exact_match=False, btn_reabrir=(By.XPATH, XPATHS_LF["btn_cliente_lov"]))
        robust_switch_to_iframe(self.driver_lf, (By.XPATH, XPATHS_LF["novo_pedido_iframe"]), (By.ID, XPATHS_LF["novo_pedido_iframe_check"]))
        wait_for_apex_ready(self.driver_lf)

        # 3. DADOS BÁSICOS
        str_tipo = str(tipo_uso).upper()
        
        if "NAN" in str_tipo or str_tipo.strip() == "":
            self.log("   ⚠️ Tipo não encontrado na planilha. Assumindo Industrialização.")
            eh_industrializacao = True
        else:
            eh_industrializacao = "CONSUMO" not in str_tipo
        
        texto_obs_pedido = OBS_FIXA_PEDIDO
        texto_obs_nf = ""

        if eh_industrializacao:
            el_fin = robust_wait_for(wait, By.ID, XPATHS_LF["campo_finalidade"])
            sel_fin = Select(el_fin)
            if sel_fin.first_selected_option.text.strip().upper() != "INDUSTRIALIZAÇÃO":
                sel_fin.select_by_visible_text("INDUSTRIALIZAÇÃO")
                wait_for_apex_ready(self.driver_lf)
                time.sleep(0.5)
            # Se for Schulz Automotiva Joinville ou Compressores, inclui a observação TTD 490
            aplica_ttd = (self.empresa_var.get() == "Automotiva" and eh_joinville) or (self.empresa_var.get() == "Compressores")
            if aplica_ttd:
                texto_obs_nf = OBS_NF_INDUSTRIALIZACAO
                self.log("   > Modo: Industrialização (Obs NF: TTD 490 Joinville)")
            else:
                self.log("   > Modo: Industrialização (Sem TTD 490)")
        else:
            el_fin = robust_wait_for(wait, By.ID, XPATHS_LF["campo_finalidade"])
            sel_fin = Select(el_fin)
            if sel_fin.first_selected_option.text.strip().upper() != "USO E CONSUMO/IMOBILIZADO":
                sel_fin.select_by_visible_text("USO E CONSUMO/IMOBILIZADO")
                # IMPORTANTE: Aguarda o APEX processar a mudança de finalidade via AJAX antes de prosseguir!
                wait_for_apex_ready(self.driver_lf)
                time.sleep(0.8)
            self.log("   > Modo: Consumo")

        # Selecionar "VENDA" no Tipo NF com espera dinâmica e retry
        wait_for_apex_ready(self.driver_lf)
        tipo_nf_selecionado = False
        for _ in range(15):
            try:
                el_tipo = robust_wait_for(wait, By.ID, XPATHS_LF["campo_tipo_nf"])
                sel_tipo = Select(el_tipo)
                if sel_tipo.first_selected_option.text.strip().upper() == "VENDA":
                    tipo_nf_selecionado = True
                    break
                sel_tipo.select_by_visible_text("VENDA")
                wait_for_apex_ready(self.driver_lf)
                tipo_nf_selecionado = True
                break
            except Exception:
                time.sleep(0.3)
        if not tipo_nf_selecionado:
            raise Exception("Não consegui selecionar o Tipo de NF 'VENDA'")

        robust_find_and_send_keys(self.driver_lf, By.ID, XPATHS_LF["campo_ordem_compra"], ordem_compra)
        wait_for_apex_ready(self.driver_lf)

        # 4. FORMA PAGAMENTO
        robust_click(self.driver_lf, By.XPATH, XPATHS_LF["btn_forma_pgto_lov"])
        self.abrir_popup_e_selecionar("CARTEIRA", exact_match=True, btn_reabrir=(By.XPATH, XPATHS_LF["btn_forma_pgto_lov"]))
        robust_switch_to_iframe(self.driver_lf, (By.XPATH, XPATHS_LF["novo_pedido_iframe"]), (By.ID, XPATHS_LF["novo_pedido_iframe_check"]))
        wait_for_apex_ready(self.driver_lf)

        # 5. CONDIÇÃO PAGTO
        robust_click(self.driver_lf, By.XPATH, XPATHS_LF["btn_condicao_pgto_lov"])
        self.abrir_popup_e_selecionar(CONDICAO_PAGTO_TEXTO, exact_match=True, btn_reabrir=(By.XPATH, XPATHS_LF["btn_condicao_pgto_lov"]))
        robust_switch_to_iframe(self.driver_lf, (By.XPATH, XPATHS_LF["novo_pedido_iframe"]), (By.ID, XPATHS_LF["novo_pedido_iframe_check"]))
        wait_for_apex_ready(self.driver_lf)

        # 6. TRANSP
        try: 
            el_frete = robust_wait_for(wait, By.ID, XPATHS_LF["campo_tipo_frete"])
            Select(el_frete).select_by_visible_text(FRETE_TIPO)
            wait_for_apex_ready(self.driver_lf, log_func=self.log)
        except: 
            pass

        dismiss_apex_alert_if_any(self.driver_lf, log_func=self.log)
        robust_click(self.driver_lf, By.XPATH, XPATHS_LF["btn_transportadora_lov"])
        self.abrir_popup_e_selecionar(CNPJ_TRANSPORTADORA, exact_match=False, btn_reabrir=(By.XPATH, XPATHS_LF["btn_transportadora_lov"]))
        robust_switch_to_iframe(self.driver_lf, (By.XPATH, XPATHS_LF["novo_pedido_iframe"]), (By.ID, XPATHS_LF["novo_pedido_iframe_check"]))
        wait_for_apex_ready(self.driver_lf, log_func=self.log)

        # 7. PEDIDO COMPLEMENTAR
        if FRETE_TIPO == "CIF" and valor_total < 400.0:
            self.log(f"   > Valor R$ {valor_total:.2f} < 400 (CIF). Marcando Pedido Complementar.")
            robust_check_element(self.driver_lf, By.XPATH, XPATHS_LF["radio_ped_complementar_sim_label"], log_func=self.log)

        # 8. OBSERVAÇÕES
        robust_find_and_send_keys(self.driver_lf, By.ID, XPATHS_LF["campo_obs_nf"], texto_obs_nf)
        robust_find_and_send_keys(self.driver_lf, By.ID, XPATHS_LF["campo_obs_pedido"], texto_obs_pedido)

        robust_click(self.driver_lf, By.XPATH, XPATHS_LF["btn_salvar_cabecalho"])
        self.driver_lf.switch_to.default_content()
        time.sleep(0.5)

    def preencher_itens(self, itens_df, ordem_compra):
        wait_lf = WebDriverWait(self.driver_lf, 25)

        for index, row in itens_df.iterrows():
            codigo_lf = str(row['cod_lf']) if pd.notna(row['cod_lf']) else ""
            qtd = row['qtd']
            preco = row['preco']
            cod_item_original = row['cod_item']

            if not codigo_lf or codigo_lf.strip() == "" or codigo_lf.lower() == "nan":
                self.log(f"   ! Item {cod_item_original} sem código LF na planilha. Pulando.")
                self.itens_nao_localizados.append(f"Pedido {ordem_compra} | Item {cod_item_original} (Sem Código LF na Planilha)")
                continue

            try:
                self.log(f"   📦 Adicionando Item: {cod_item_original} (LF: {codigo_lf}) | Qtd: {qtd} | Preço: R$ {preco:.2f}")
                robust_click(self.driver_lf, By.XPATH, XPATHS_LF["btn_adicionar_item"])
                robust_switch_to_iframe(self.driver_lf, (By.XPATH, XPATHS_LF["item_popup_iframe"]), (By.ID, XPATHS_LF["item_popup_iframe_check"]))
                
                robust_click(self.driver_lf, By.ID, XPATHS_LF["item_popup_sku_btn"])
                sucesso_busca = self.abrir_popup_e_selecionar(
                    codigo_lf, 
                    exact_match=False, 
                    ignorar_erro=True,
                    btn_reabrir=(By.ID, XPATHS_LF["item_popup_sku_btn"]),
                    iframe_info=((By.XPATH, XPATHS_LF["item_popup_iframe"]), (By.ID, XPATHS_LF["item_popup_iframe_check"]))
                ) 
                
                if not sucesso_busca:
                    self.log(f"   ! LF {codigo_lf} não encontrado. Pulando item.")
                    self.itens_nao_localizados.append(f"Pedido {ordem_compra} | Item {cod_item_original} | LF {codigo_lf} (Não encontrado no portal)")
                    self.driver_lf.switch_to.default_content()
                    try: 
                        self.driver_lf.find_element(By.CSS_SELECTOR, ".ui-dialog-titlebar-close").click()
                    except: 
                        pass
                    continue
                
                robust_switch_to_iframe(self.driver_lf, (By.XPATH, XPATHS_LF["item_popup_iframe"]), (By.ID, XPATHS_LF["item_popup_iframe_check"]))
                
                robust_find_and_send_keys(self.driver_lf, By.ID, XPATHS_LF["item_popup_ordem_compra"], ordem_compra)
                robust_find_and_send_keys(self.driver_lf, By.ID, XPATHS_LF["item_popup_qtd"], str(qtd).replace('.', ','))
                robust_find_and_send_keys(self.driver_lf, By.ID, XPATHS_LF["item_popup_preco"], str(preco).replace('.', ','))
                
                robust_click(self.driver_lf, By.ID, XPATHS_LF["item_popup_btn_salvar"])
                self.driver_lf.switch_to.default_content()
                wait_lf.until(EC.invisibility_of_element_located((By.XPATH, XPATHS_LF["item_popup_iframe"])))

            except Exception as e:
                self.log(f"   ! Erro ao processar item {cod_item_original}: {e}")
                self.itens_nao_localizados.append(f"Pedido {ordem_compra} | Item {cod_item_original} | Erro: {str(e)}")
                try: 
                    self.driver_lf.find_element(By.CSS_SELECTOR, ".ui-dialog-titlebar-close").click()
                except: 
                    pass
                self.driver_lf.switch_to.default_content()
                try: 
                    self.driver_lf.find_element(By.XPATH, "//button[@title='Close']").click()
                except: 
                    pass

if __name__ == "__main__":
    root = tk.Tk()
    app = AutomacaoSchulzApp(root)
    root.mainloop()